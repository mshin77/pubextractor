import email
import io
import shutil
import subprocess
import urllib.request
import urllib.response

import pytest
from docx import Document
from docx.oxml import parse_xml

from pubextractor import providers
from pubextractor.docx_reader import read_docx
from pubextractor.licensing import is_local_endpoint
from pubextractor.providers import _post, _RefuseRedirect, load_env, make_responder

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


@pytest.fixture
def redirecting_opener(monkeypatch):
    seen = []

    class Redirecting(urllib.request.HTTPHandler):
        def http_open(self, req):
            seen.append(req.get_header('Authorization'))
            headers = email.message_from_string('Location: http://127.0.0.1:1/next?token=zzz\n\n')
            reply = urllib.response.addinfourl(io.BytesIO(b''), headers, req.full_url, 307)
            reply.msg = 'Temporary Redirect'
            return reply

    opener = urllib.request.build_opener(Redirecting, _RefuseRedirect)
    monkeypatch.setattr(providers, '_OPENER', opener)
    return seen


def test_post_refuses_redirect(redirecting_opener):
    with pytest.raises(RuntimeError, match='redirect') as caught:
        _post('http://127.0.0.1:9/v1', {}, {'Authorization': 'Bearer k'}, 5)
    assert 'zzz' not in str(caught.value)
    assert len(redirecting_opener) == 1


def test_keyed_provider_with_base_url_needs_key_env():
    with pytest.raises(ValueError, match='key_env'):
        make_responder('openai', 'm', base_url='http://127.0.0.1:8000/v1',
                       environ={'OPENAI_API_KEY': 'k'}, post=lambda *a: {})


def test_key_over_plain_http_to_remote_refused():
    with pytest.raises(ValueError, match='http'):
        make_responder('custom', 'm', base_url='http://example.org/v1', key_env='K',
                       environ={'K': 'k'}, post=lambda *a: {})


def test_key_over_https_remote_allowed():
    make_responder('custom', 'm', base_url='https://example.org/v1', key_env='K',
                   environ={'K': 'k'}, post=lambda *a: {})


class TestEnvParsing:
    def load(self, tmp_path, text):
        env_file = tmp_path / '.env'
        env_file.write_text(text)
        environ = {}
        load_env(env_file, environ)
        return environ

    def test_inline_comment_stripped(self, tmp_path):
        assert self.load(tmp_path, 'A=abc # note\n')['A'] == 'abc'

    def test_hash_inside_value_kept(self, tmp_path):
        assert self.load(tmp_path, 'A=ab#c\n')['A'] == 'ab#c'

    def test_quoted_value_keeps_hash(self, tmp_path):
        assert self.load(tmp_path, 'A="a # b"\n')['A'] == 'a # b'

    def test_mismatched_quotes_kept_literal(self, tmp_path):
        assert self.load(tmp_path, 'C="x"y\n')['C'] == '"x"y'

    def test_empty_value_skipped(self, tmp_path):
        env_file = tmp_path / '.env'
        env_file.write_text('B=\nB=real\n')
        environ = {}
        assert load_env(env_file, environ) == ['B']
        assert environ == {'B': 'real'}

    def test_empty_value_not_reported(self, tmp_path):
        env_file = tmp_path / '.env'
        env_file.write_text('B=\n')
        assert load_env(env_file, {}) == []


def test_warns_when_env_is_tracked(tmp_path, monkeypatch):
    if shutil.which('git') is None:
        pytest.skip('git is not installed')
    monkeypatch.setenv('GIT_CONFIG_GLOBAL', '')
    monkeypatch.setenv('GIT_CONFIG_NOSYSTEM', '1')
    subprocess.run(['git', 'init', '-q', str(tmp_path)], check=True)
    env_file = tmp_path / '.env'
    env_file.write_text('A=1\n')
    subprocess.run(['git', 'add', '-f', '.env'], cwd=tmp_path, check=True)
    (tmp_path / '.gitignore').write_text('.env\n')
    with pytest.warns(UserWarning, match='tracked'):
        load_env(env_file, {})


@pytest.mark.parametrize('url', ['http://localhost\t/', 'http://localhost /', 'http://local\nhost/',
                                 'http://127.0.0.1\x00/'])
def test_endpoint_with_whitespace_or_control_not_local(url):
    assert not is_local_endpoint(url)


def make_sdt_document(path):
    document = Document()
    document.add_paragraph('before')
    body = document.element.body
    sdt = parse_xml(
        f'<w:sdt {W}><w:sdtPr/><w:sdtContent><w:p><w:r><w:t>inside control</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>')
    body.insert(len(body) - 1, sdt)
    document.add_paragraph('after')
    document.save(path)


def test_sdt_content_included_in_order(tmp_path):
    path = tmp_path / 'a.docx'
    make_sdt_document(path)
    texts = [b['text'] for b in read_docx(path) if b['kind'] == 'paragraph']
    assert texts == ['before', 'inside control', 'after']


def test_merged_cell_text_once(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'span'
    table.cell(0, 2).text = 'c'
    for i, text in enumerate(['a', 'b', 'c2']):
        table.cell(1, i).text = text
    path = tmp_path / 'm.docx'
    document.save(path)
    rows = [b for b in read_docx(path) if b['kind'] == 'table'][0]['rows']
    assert rows == [['span', '', 'c'], ['a', 'b', 'c2']]


def test_doc_input_gives_clear_error(tmp_path):
    path = tmp_path / 'old.doc'
    path.write_bytes(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 64)
    with pytest.raises(ValueError, match='only .docx is supported'):
        read_docx(path)
