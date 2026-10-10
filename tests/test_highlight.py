import pytest


class TestPatternFromTerms:
    def build(self, terms):
        from pubextractor.highlight import pattern_from_terms

        return pattern_from_terms(terms)

    def test_a_plain_word_also_finds_its_plural(self):
        import re

        pattern = self.build('middle school, class')
        assert re.search(pattern, 'two middle schools took part', re.I)
        assert re.search(pattern, 'four classes', re.I)
        assert not re.search(pattern, 'middle schooling', re.I)

    def test_truncation_matches_the_family(self):
        import re

        pattern = self.build('dyslex*')
        assert re.search(pattern, 'a study of dyslexia', re.I)
        assert re.search(pattern, 'dyslexic readers', re.I)

    def test_truncation_still_respects_word_starts(self):
        import re

        assert not re.search(self.build('dyslex*'), 'predyslexial notions', re.I)

    def test_a_phrase_matches_across_any_whitespace(self):
        import re

        pattern = self.build('higher education')
        assert re.search(pattern, 'in higher education today', re.I)
        assert re.search(pattern, 'in higher  education today', re.I)
        assert re.search(pattern, 'in higher\neducation today', re.I)

    def test_terms_are_literal_not_patterns(self):
        import re

        pattern = self.build('e.g.')
        assert re.search(pattern, 'see e.g. the review', re.I)
        assert not re.search(pattern, 'see eXgY the review', re.I)

    def test_several_terms_become_alternatives(self):
        import re

        pattern = self.build('university, college, higher education')
        for text in ('at university', 'a college course', 'in higher education'):
            assert re.search(pattern, text, re.I), text

    def test_blank_terms_are_dropped(self):
        assert self.build('dyslexia, , ') == self.build('dyslexia')

    def test_no_terms_gives_no_pattern(self):
        assert self.build('') == ''
        assert self.build([]) == ''

    def test_a_sequence_works_as_well_as_a_string(self):
        assert self.build(['dyslexia', 'college']) == self.build('dyslexia, college')


    def test_a_star_inside_a_phrase_truncates_that_word(self):
        import re

        pattern = self.build('applied behavior* analys*')
        assert re.search(pattern, 'applied behavior analysis', re.I)
        assert re.search(pattern, 'applied behavioral analyses', re.I)

    def test_a_star_inside_a_word_truncates_there(self):
        import re

        pattern = self.build('speech*to*text')
        assert re.search(pattern, 'speech-to-text output', re.I)

    def test_a_hyphen_stands_in_for_a_space(self):
        import re

        pattern = self.build('single* case* design')
        assert re.search(pattern, 'a single-case design', re.I)
        assert re.search(pattern, 'a single case design', re.I)

    def test_a_star_never_escapes_as_a_literal(self):
        pattern = self.build('multiple* baseline* design')
        assert '\\*' not in pattern

    def test_the_result_drives_a_pandas_search(self):
        import pandas as pd

        titles = pd.Series(['Dyslexia at university', 'Unrelated work'])
        hits = titles.str.contains(self.build('dyslex*'), case=False, regex=True)
        assert hits.tolist() == [True, False]


class TestHighlightEvidence:
    def workbook(self, tmp_path, evidence='THE STUDY USED A MULTIPLE BASELINE DESIGN.'):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        for column, name in enumerate(['title', 'abstract', 'design_evidence'], start=1):
            sheet.cell(1, column, name)
        sheet.cell(2, 1, 'A multiple baseline study')
        sheet.cell(2, 2, 'Reading gains under a multiple baseline design.')
        sheet.cell(2, 3, evidence)
        path = tmp_path / 'screened.xlsx'
        book.save(path)
        return path

    def patterns(self):
        return {'design': r'multiple\s+baseline'}

    def test_it_marks_the_matched_term(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        marked = highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])
        assert marked == 3

        sheet = load_workbook(path, rich_text=True).active
        assert type(sheet.cell(2, 1).value).__name__ == 'CellRichText'

    def test_text_survives_the_marking(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])
        sheet = load_workbook(path, rich_text=True).active
        joined = ''.join(str(part) for part in sheet.cell(2, 2).value)
        assert joined == 'Reading gains under a multiple baseline design.'

    def test_a_cell_without_a_match_is_left_alone(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path, evidence='No design named here.')
        highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])
        sheet = load_workbook(path, rich_text=True).active
        assert sheet.cell(2, 3).value == 'No design named here.'

    def test_marking_twice_changes_nothing(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        first = highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])
        second = highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])
        assert first == second

        sheet = load_workbook(path, rich_text=True).active
        joined = ''.join(str(part) for part in sheet.cell(2, 1).value)
        assert joined == 'A multiple baseline study'


    def test_spacing_around_a_marked_term_is_kept(self, tmp_path):
        import re
        import zipfile

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])

        book = zipfile.ZipFile(path)
        sheet = [n for n in book.namelist() if n.startswith('xl/worksheets/')][0]
        xml = book.read(sheet).decode('utf-8')
        assert not re.findall(r'<t>(\s[^<]*|[^<]*\s)</t>', xml)

    def test_the_workbook_stays_readable(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        highlight_evidence(path, self.patterns(), fields=['title', 'abstract'])

        sheet = load_workbook(path, rich_text=True).active
        joined = ''.join(str(part) for part in sheet.cell(2, 1).value)
        assert joined == 'A multiple baseline study'

    def test_a_missing_field_is_skipped(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        marked = highlight_evidence(path, self.patterns(), fields=['title', 'nothing_here'])
        assert marked == 2


class TestHighlightLongLayout:
    def workbook(self, tmp_path):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        for column, name in enumerate(['study_id', 'field', 'value', 'quote', 'pdf_page'],
                                      start=1):
            sheet.cell(1, column, name)
        sheet.append(['s1', 'grade', '7', 'Seventh grade students in a multiple baseline design.', 4])
        sheet.append(['s1', 'design', 'MBD', 'Seventh grade students in a multiple baseline design.', 4])
        sheet.append(['s1', 'unknown', '', 'Seventh grade students.', 5])
        path = tmp_path / 'long.xlsx'
        book.save(path)
        return path

    def patterns(self):
        return {'grade': r'seventh\s+grade', 'design': r'multiple\s+baseline'}

    def marked_text(self, cell):
        return [str(part.text) for part in cell.value if hasattr(part, 'font')]

    def test_each_row_shows_only_its_own_field(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        marked = highlight_evidence(path, self.patterns(), key_column='field')
        assert marked == 2

        sheet = load_workbook(path, rich_text=True).active
        assert self.marked_text(sheet.cell(2, 4)) == ['Seventh grade']
        assert self.marked_text(sheet.cell(3, 4)) == ['multiple baseline']

    def test_row_with_unknown_field_is_left_alone(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        highlight_evidence(path, self.patterns(), key_column='field')
        sheet = load_workbook(path, rich_text=True).active
        assert sheet.cell(4, 4).value == 'Seventh grade students.'
        assert sheet.cell(2, 5).value == 4

    def test_color_follows_the_field(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.workbook(tmp_path)
        highlight_evidence(path, self.patterns(), key_column='field',
                           colours={'grade': 'FF0000FF', 'design': 'FF00AA00'})
        sheet = load_workbook(path, rich_text=True).active
        fonts = [part.font.color.rgb for part in sheet.cell(2, 4).value if hasattr(part, 'font')]
        assert fonts == ['FF0000FF']


class TestHighlightHeaderGuards:
    def build(self, tmp_path, header, rows):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        path = tmp_path / 'guards.xlsx'
        book.save(path)
        return path

    def test_blank_header_cell_keeps_the_wide_layout(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['title', None, 'abstract'],
                          [['multiple baseline study', 'x', 'a multiple baseline design']])
        assert highlight_evidence(path, {'design': r'multiple\s+baseline'}) == 2

    def test_missing_key_column_raises(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['field', 'quote'], [['grade', 'seventh grade']])
        with pytest.raises(KeyError, match='criterion'):
            highlight_evidence(path, {'grade': 'seventh'}, key_column='criterion')

    def test_key_with_spacing_or_number_still_matches(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['field', 'quote'],
                          [['grade ', 'seventh grade'], [1, 'item 1 here']])
        assert highlight_evidence(path, {'grade': 'seventh', '1': 'item'},
                                  key_column='field') == 2

    def test_empty_pattern_is_skipped(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['title'], [['any title']])
        assert highlight_evidence(path, {'blank': ''}) == 0


class TestHighlightNameAndPatternGuards:
    def build(self, tmp_path, header, rows):
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        path = tmp_path / 'names.xlsx'
        book.save(path)
        return path

    def test_color_follows_trimmed_name(self, tmp_path):
        from openpyxl import load_workbook

        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['title'], [['seventh grade']])
        highlight_evidence(path, {' grade ': 'grade'}, colours={'grade': 'FF00FF00'})
        cell = load_workbook(path, rich_text=True).active.cell(2, 1).value
        assert [part.font.color.rgb for part in cell if hasattr(part, 'font')] == ['FF00FF00']

    def test_names_that_trim_alike_are_rejected(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['title'], [['x']])
        with pytest.raises(ValueError, match='grade'):
            highlight_evidence(path, {'grade': 'a', 'grade ': 'b'})

    def test_header_with_spaces_is_found(self, tmp_path):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, [' title '], [['multiple baseline study']])
        assert highlight_evidence(path, {'design': 'multiple baseline'}) == 1

    @pytest.mark.parametrize('pattern', ['x?', ' '])
    def test_empty_or_blank_matches_mark_nothing(self, tmp_path, pattern):
        from pubextractor.highlight import highlight_evidence

        path = self.build(tmp_path, ['title'], [['a b c']])
        assert highlight_evidence(path, {'blank': pattern}) == 0


class TestPlurals:
    @pytest.mark.parametrize('term, text', [
        ('study', 'three studies'), ('study', 'one study'), ('analysis', 'two analyses'),
        ('hypothesis', 'many hypotheses'), ('child', 'the children'), ('person', 'people'),
        ('man', 'men'), ('woman', 'women'), ('datum', 'the data'),
        ('criterion', 'all criteria'), ('phenomenon', 'rare phenomena'),
        ('day', 'two days'), ('class', 'four classes')])
    def test_plural_forms_are_found(self, term, text):
        import re

        from pubextractor.highlight import pattern_from_terms

        assert re.search(pattern_from_terms(term), text, re.I)

    @pytest.mark.parametrize('term, text', [
        ('class', 'a classic'), ('study', 'studying'), ('day', 'daies'), ('man', 'mens'),
        ('analysis', 'analysed')])
    def test_plural_rules_do_not_widen_the_match(self, term, text):
        import re

        from pubextractor.highlight import pattern_from_terms

        assert not re.search(pattern_from_terms(term), text, re.I)

    def test_star_keeps_truncation_without_plural_rules(self):
        import re

        from pubextractor.highlight import pattern_from_terms

        pattern = pattern_from_terms('stud*')
        assert re.search(pattern, 'studies and studying', re.I)
