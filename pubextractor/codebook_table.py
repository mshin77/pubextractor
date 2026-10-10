"""Codebook as one long sheet, one row per term, for editing in Excel or a CSV editor."""
import csv
import logging
import os
from pathlib import Path

from .codebook import _problems, _text, load_codebook

logger = logging.getLogger(__name__)

COLUMNS = ('field', 'type', 'value', 'term', 'definition', 'query', 'stat', 'range_min',
           'range_max', 'selected', 'color')
REQUIRED = ('field', 'type')
CATEGORY = ('category', 'multi_category')
TRUE = {'true', 'yes', '1'}
FALSE = {'false', 'no', '0'}


def _openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise ImportError('Excel files require openpyxl; reinstall pubextractor') from None
    return openpyxl


def _cell(value):
    if isinstance(value, str):
        value = value.strip()
    return None if value is None or value == '' else value


def _lowered(column, value):
    return value.lower() if column == 'type' and isinstance(value, str) else value


def _suffix(path):
    suffix = Path(path).suffix.lower()
    if suffix not in ('.xlsx', '.csv'):
        raise ValueError(f"{os.fspath(path)!r}: a codebook table must be .xlsx or .csv")
    return suffix


def _read_cells(path):
    if _suffix(path) == '.csv':
        with open(path, encoding='utf-8-sig', newline='') as handle:
            return [tuple(row) for row in csv.reader(handle)]
    openpyxl = _openpyxl()
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(book.worksheets[0].iter_rows(values_only=True))
        finally:
            book.close()
    except Exception as error:
        raise ValueError(f"could not read {Path(path).name}: {error}") from error


def _number(value):
    if isinstance(value, bool):
        raise ValueError
    number = value if isinstance(value, (int, float)) else float(value)
    return int(number) if float(number).is_integer() else float(number)


def _selected(value):
    text = str(value).strip().lower()
    if value is None or isinstance(value, bool):
        return value is not False
    if text in TRUE | FALSE:
        return text in TRUE
    raise ValueError


def _parsed(number, row):
    found, problems = {}, []
    for column in ('range_min', 'range_max'):
        try:
            found[column] = None if row.get(column) is None else _number(row[column])
        except ValueError:
            problems.append(f"row {number}: {column} must be a number, not {row[column]!r}")
    try:
        found['selected'] = _selected(row.get('selected'))
    except ValueError:
        problems.append(f"row {number}: selected must be TRUE or FALSE, not {row['selected']!r}")
    return found, problems


def _row_problems(number, name, kind, row, first_kind, first_row):
    category = kind in CATEGORY
    checks = [
        (kind != first_kind, f"{name}: type {kind} differs from row {first_row}"),
        (category and row.get('value') is None, f"{name}: value is empty"),
        (not category and row.get('value') is not None,
         f"{name}: value must be empty for {kind} fields"),
        (not category and row.get('term') is not None,
         f"{name}: term must be empty for {kind} fields"),
        ((row.get('range_min') is None) != (row.get('range_max') is None),
         'range needs both range_min and range_max'),
    ]
    return [f"row {number}: {message}" for bad, message in checks if bad]


def _fields(rows):
    fields, first, problems, seen = {}, {}, [], {}
    for number, row in rows:
        if row.get('field') is None:
            problems.append(f"row {number}: field is empty")
            continue
        if row.get('type') is None:
            problems.append(f"row {number}: type is empty")
            continue
        name, kind = _text(row['field']), row['type']
        identity = tuple(None if row.get(column) is None else _text(row[column])
                         for column in ('field', 'value', 'term'))
        if identity in seen:
            logger.warning('row %d duplicates row %d (field, value, term) and is dropped',
                           number, seen[identity])
            continue
        seen[identity] = number
        first.setdefault(name, number)
        field = fields.setdefault(name, {'type': kind, 'values': [], 'terms': {}})
        parsed, bad = _parsed(number, row)
        problems += bad + _row_problems(number, name, kind, row, field['type'], first[name])
        value = None if row.get('value') is None else _text(row['value'])
        if field['type'] in CATEGORY and value is not None and value not in field['values']:
            field['values'].append(value)
        if value is not None and row.get('term') is not None and parsed.get('selected', True):
            field['terms'].setdefault(value, []).append(_text(row['term']))
        span = [parsed.get('range_min'), parsed.get('range_max')]
        if None not in span:
            field.setdefault('range', span)
        for column in ('definition', 'query', 'stat', 'color'):
            if row.get(column) is not None:
                field.setdefault(column, _text(row[column]))
    problems += [f"row {first[name]}: {message}"
                 for name, field in fields.items() for message in _problems(name, field)]
    return fields, problems


def read_codebook_table(path):
    """Codebook from a long .xlsx or .csv sheet, one row per term.

    Columns are ``field, type, value, term, definition, query, stat,
    range_min, range_max, selected, color``; only ``field`` and ``type`` are
    required. A category value with no terms takes one row with an empty term;
    a number, integer, or text field takes one row with ``value`` and ``term``
    empty. ``selected`` FALSE drops that term. The first sheet is read.

    Raises ``ValueError`` naming the sheet row of every problem at once.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"{os.fspath(path)!r} is not a file")
    cells = _read_cells(path)
    header = [str(name).strip().lower() if name is not None else '' for name in
              (cells[0] if cells else ())]
    missing = [column for column in REQUIRED if column not in header]
    if missing:
        raise ValueError(f"Codebook table is missing columns: {', '.join(missing)}")
    rows = [(number, {column: _lowered(column, _cell(value))
                      for column, value in zip(header, line) if column})
            for number, line in enumerate(cells[1:], start=2)
            if any(_cell(value) is not None for value in line)]
    fields, problems = _fields(rows)
    if problems:
        raise ValueError('Codebook table problems: ' + '; '.join(problems))
    return load_codebook(fields)


def _rows(name, field):
    span = field.get('range') or ('', '')
    shared = {'field': name, 'type': field['type'], 'definition': field.get('definition') or '',
              'query': field.get('query') or '', 'stat': field.get('stat') or '',
              'range_min': span[0], 'range_max': span[1], 'selected': True,
              'color': field.get('color') or ''}
    pairs = ([(value, term) for value in field['values']
              for term in field['terms'].get(value) or ['']]
             if field['type'] in CATEGORY else [('', '')])
    return [{**shared, 'value': value, 'term': term} for value, term in pairs]


def write_codebook_table(codebook, path):
    """Write a codebook as a long .xlsx or .csv sheet that ``read_codebook_table`` reads back.

    Returns the path written.
    """
    path = Path(path)
    suffix = _suffix(path)
    rows = [row for name, field in load_codebook(codebook).items() for row in _rows(name, field)]
    if suffix == '.csv':
        with open(path, 'w', encoding='utf-8', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows({**row, 'selected': 'TRUE'} for row in rows)
        return path
    book = _openpyxl().Workbook()
    sheet = book.active
    sheet.title = 'codebook'
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append([None if row[column] == '' else row[column] for column in COLUMNS])
    sheet.freeze_panes = 'A2'
    book.save(path)
    return path
