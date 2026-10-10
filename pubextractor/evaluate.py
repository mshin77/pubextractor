"""Agreement between extracted values and a second coder, per field and per graph point."""
import numpy as np

from .codebook_table import _openpyxl
from .output import COLUMNS, to_long

AGREEMENT_COLUMNS = ('field', 'n_pairs', 'agreement', 'kappa', 'icc', 'mae', 'bias', 'lower',
                     'upper', 'hallucinations', 'omissions')
SHEET_COLUMNS = (*COLUMNS, 'coded_value')
EMPTY = (None, '', 'not_found')


def kappa(a, b):
    """Cohen's kappa for two raters' labels; None with fewer than two pairs or chance agreement 1."""
    if len(a) < 2:
        return None
    labels, codes = np.unique(np.asarray([*a, *b], dtype=str), return_inverse=True)
    table = np.zeros((len(labels), len(labels)))
    np.add.at(table, (codes[:len(a)], codes[len(a):]), 1 / len(a))
    chance = table.sum(axis=1) @ table.sum(axis=0)
    return None if np.isclose(chance, 1) else float((np.trace(table) - chance) / (1 - chance))


def icc(ratings):
    """ICC(2,1) from an n targets by k raters array; None with fewer than two targets.

    Two-way random effects, absolute agreement, single rater: ICC(2,1) of Shrout
    and Fleiss (1979), ICC(A,1) of McGraw and Wong (1996), from the mean squares
    (MSR - MSE) / (MSR + (k - 1) MSE + k (MSC - MSE) / n).
    """
    y = np.asarray(ratings, dtype=float)
    n, k = y.shape
    if n < 2:
        return None
    rows, cols, grand = y.mean(axis=1), y.mean(axis=0), y.mean()
    msr = k * ((rows - grand) ** 2).sum() / (n - 1)
    msc = n * ((cols - grand) ** 2).sum() / (k - 1)
    mse = ((y - rows[:, None] - cols + grand) ** 2).sum() / ((n - 1) * (k - 1))
    denominator = msr + (k - 1) * mse + k * (msc - mse) / n
    return float((msr - mse) / denominator) if denominator else None


def _number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def _label(value):
    return 'not_found' if value is None else str(value).strip().casefold()


def evaluate(extracted, coded, key=('paper_id', 'field'), tolerance=None):
    """Per-field agreement between extracted records and a human-coded sheet.

    Both are long records; the coder's value is in ``value``. Records pair on
    ``key``, the first record per key counting. A pair agrees when both have no
    value, when categories match after case and spaces, or when numbers are
    equal after rounding to 2 decimals (within ``tolerance`` when given). A
    field is numeric when every extracted value in it reads as a number (a coded
    value that is not a number then counts as a disagreement): it gets ICC(2,1),
    MAE, and Bland-Altman bias with 95% limits (bias +/- 1.96 SD of the
    extracted minus coded differences); any other field gets Cohen's kappa, with
    no value as its own category. A hallucination is an extracted value where
    the coder found nothing (coded ``not_found`` or blank); an omission is a
    value the coder found where extraction gave none (no value or ``not_found``).
    Returns one row per field in ``AGREEMENT_COLUMNS``; agreement is a percent.
    """
    found, truth = {}, {}
    [found.setdefault(tuple(r.get(c) for c in key),
                      None if r.get('status') == 'not_found' or r.get('value') in EMPTY
                      else r['value']) for r in extracted]
    [truth.setdefault(tuple(r.get(c) for c in key),
                      None if r.get('value') in EMPTY else r['value']) for r in coded]
    shared = [k for k in found if k in truth]
    at = key.index('field')
    rows = []
    for name in dict.fromkeys(k[at] for k in shared):
        pairs = [(found[k], truth[k]) for k in shared if k[at] == name]
        given = [a for a, _ in pairs if a is not None]
        numeric = bool(given) and not np.isnan([_number(a) for a in given]).any()
        numbers = np.asarray([[_number(v) for v in pair] for pair in pairs if None not in pair]
                             ).reshape(-1, 2)
        numbers = numbers[~np.isnan(numbers).any(axis=1)]
        difference = numbers[:, 0] - numbers[:, 1]
        close = (np.abs(difference) <= tolerance if tolerance is not None
                 else np.round(numbers[:, 0], 2) == np.round(numbers[:, 1], 2))
        labels = np.asarray([[_label(v) for v in pair] for pair in pairs])
        agree = (sum(pair == (None, None) for pair in pairs) + close.sum() if numeric
                 else (labels[:, 0] == labels[:, 1]).sum())
        spread = 1.96 * difference.std(ddof=1) if len(difference) > 1 else None
        bias = float(difference.mean()) if numeric else None
        stats = ({'icc': icc(numbers), 'mae': float(np.abs(difference).mean()), 'bias': bias,
                  'lower': None if spread is None else bias - spread,
                  'upper': None if spread is None else bias + spread}
                 if numeric else {'kappa': kappa(labels[:, 0], labels[:, 1])})
        rows.append({**dict.fromkeys(AGREEMENT_COLUMNS), 'field': name, 'n_pairs': len(pairs),
                     'agreement': 100 * float(agree) / len(pairs), **stats,
                     'hallucinations': sum(a is not None and b is None for a, b in pairs),
                     'omissions': sum(a is None and b is not None for a, b in pairs)})
    return rows


def _points(rows):
    return {(str(r['case']), round(float(r['session']))): float(r['outcome']) for r in rows}


def graph_agreement(extracted_rows, manual_rows, y_range=None):
    """Point-level error of extracted graph rows against a manual digitizer.

    Points pair on ``case`` and ``session`` (rounded to a whole session).
    Returns ``n_pairs``, ``mae`` in outcome units, ``mae_percent`` of
    ``y_range`` (the y-axis top minus bottom; the span of the manual outcomes
    when not given), ``missing`` manual points with no extracted point, and
    ``extra`` extracted points with no manual point.
    """
    found, manual = _points(extracted_rows), _points(manual_rows)
    shared = [k for k in manual if k in found]
    mae = float(np.mean([abs(found[k] - manual[k]) for k in shared])) if shared else None
    span = y_range or (np.ptp(list(manual.values())) if manual else None)
    return {'n_pairs': len(shared), 'mae': mae,
            'mae_percent': 100 * mae / span if mae is not None and span else None,
            'missing': len(manual) - len(shared), 'extra': len(found) - len(shared)}


def write_coding_sheet(records, path):
    """Write records as an .xlsx for a second coder: the long columns and an empty ``coded_value``.

    Returns the path written.
    """
    openpyxl = _openpyxl()
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'coding'
    sheet.append(list(SHEET_COLUMNS))
    [sheet.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str)
                   else v if v is None or isinstance(v, (int, float)) else str(v)
                   for v in row.values()]) for row in to_long(records)]
    sheet.freeze_panes = 'A2'
    book.save(path)
    return path


def read_coding_sheet(path):
    """Coded records from a filled coding sheet, with the coder's ``coded_value`` as ``value``."""
    book = _openpyxl().load_workbook(path, read_only=True, data_only=True)
    lines = list(book.worksheets[0].iter_rows(values_only=True))
    book.close()
    rows = [dict(zip(lines[0], line)) for line in lines[1:] if any(v is not None for v in line)]
    return [{**{c: v for c, v in row.items() if c != 'coded_value'}, 'value': row.get('coded_value')}
            for row in rows]
